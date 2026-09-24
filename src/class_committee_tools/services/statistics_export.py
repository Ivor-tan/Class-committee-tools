"""将当前学生统计数据导出为 Excel。"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from class_committee_tools.models import ActivityRequirement, Student
from class_committee_tools.repositories import DatabaseRepository

EXPORT_FIELDS = {
    "name": "姓名",
    "student_number": "学号",
    "status": "提交状态",
    "progress": "材料进度",
    "file_count": "文件数量",
    "updated_at": "最后更新时间",
    "requirement_details": "各材料提交情况",
}


def export_student_statistics(
    path: Path,
    students: list[Student],
    requirements: list[ActivityRequirement],
    repository: DatabaseRepository,
    selected_fields: set[str],
) -> Path:
    """导出学生数据，不包含任何材料文件内容。"""
    if not selected_fields:
        raise ValueError("请至少选择一项要导出的数据。")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl 依赖，无法导出 Excel。") from exc

    path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生提交统计"

    headers: list[str] = []
    for field, title in EXPORT_FIELDS.items():
        if field == "requirement_details":
            if field in selected_fields:
                headers.extend(item.name for item in requirements)
        elif field in selected_fields:
            headers.append(title)
    sheet.append(headers)

    for student in students:
        submitted_requirement_ids = {
            material.requirement_id
            for material in repository.list_materials(student.id)
            if material.requirement_id is not None
        }
        row: list[str | int] = []
        for field in EXPORT_FIELDS:
            if field not in selected_fields:
                continue
            if field == "name":
                row.append(student.name)
            elif field == "student_number":
                row.append(student.student_number or "")
            elif field == "status":
                row.append("已提交" if student.submitted else "未提交")
            elif field == "progress":
                row.append(
                    f"必交 {student.fulfilled_count}/{student.required_count} 项"
                    if student.requirement_count
                    else f"{student.file_count} 个文件"
                )
            elif field == "file_count":
                row.append(student.file_count)
            elif field == "updated_at":
                row.append(student.updated_at.strftime("%Y-%m-%d %H:%M:%S"))
            elif field == "requirement_details":
                row.extend(
                    "已提交" if item.id in submitted_requirement_ids else "未提交"
                    for item in requirements
                )
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, column_cells in enumerate(sheet.columns, start=1):
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 3, 12), 40
        )

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix="student_statistics_", suffix=".xlsx", dir=path.parent, delete=False
        ) as handle:
            temporary_path = Path(handle.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
        temporary_path = None
    finally:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
    return path
