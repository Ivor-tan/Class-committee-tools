"""Excel 学生名单解析。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

NAME_HEADERS = ("姓名", "学生姓名", "名字")
NUMBER_HEADERS = ("学号", "学生学号", "编号")


@dataclass(frozen=True, slots=True)
class ImportPreview:
    entries: list[tuple[str, str | None]]
    ignored_empty: int
    ignored_duplicates: int
    name_header: str
    number_header: str | None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_student_workbook(path: Path) -> ImportPreview:
    """读取首个工作表并识别姓名及可选学号列。"""
    if path.suffix.lower() != ".xlsx":
        raise ValueError("当前版本仅支持 .xlsx 文件。")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl 依赖，无法读取 Excel。") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 Excel 文件：{exc}") from exc

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        if not header_values:
            raise ValueError("Excel 中没有可读取的数据。")
        headers = [_cell_text(value) for value in header_values]
        name_index = next((headers.index(item) for item in NAME_HEADERS if item in headers), None)
        if name_index is None:
            raise ValueError("未找到“姓名”“学生姓名”或“名字”列。")
        number_index = next(
            (headers.index(item) for item in NUMBER_HEADERS if item in headers), None
        )

        entries: list[tuple[str, str | None]] = []
        seen: set[tuple[str, str]] = set()
        empty_count = 0
        duplicate_count = 0
        for row in rows:
            name = _cell_text(row[name_index]) if name_index < len(row) else ""
            if not name:
                empty_count += 1
                continue
            number_text = (
                _cell_text(row[number_index])
                if number_index is not None and number_index < len(row)
                else ""
            )
            key = (name, number_text)
            if key in seen:
                duplicate_count += 1
                continue
            seen.add(key)
            entries.append((name, number_text or None))
        if not entries:
            raise ValueError("没有找到有效的学生姓名。")
        return ImportPreview(
            entries,
            empty_count,
            duplicate_count,
            headers[name_index],
            headers[number_index] if number_index is not None else None,
        )
    finally:
        workbook.close()
