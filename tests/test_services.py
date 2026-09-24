from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from class_committee_tools.repositories import DatabaseRepository
from class_committee_tools.services import (
    EXPORT_FIELDS,
    ActivityService,
    ExportService,
    FileService,
    export_student_statistics,
)
from class_committee_tools.utils import sanitize_filename


@pytest.fixture
def repository(tmp_path: Path) -> DatabaseRepository:
    return DatabaseRepository(tmp_path / "app.db")


def test_activity_validation_and_persistence(repository: DatabaseRepository) -> None:
    service = ActivityService(repository)
    with pytest.raises(ValueError):
        service.create("  ")
    with pytest.raises(ValueError):
        service.create("非法/名称")

    activity = service.create("奖学金材料")

    assert repository.get_activity(activity.id) is not None
    assert repository.get_activity(activity.id).name == "奖学金材料"
    assert repository.get_activity(activity.id).status == "active"


def test_activity_can_be_renamed_completed_and_deleted(
    repository: DatabaseRepository,
) -> None:
    service = ActivityService(repository)
    activity = service.create("待修改活动")
    repository.add_students(activity.id, [("张三", None)])

    service.rename(activity.id, "新活动名称")
    service.set_completed(activity.id, True)

    updated = repository.get_activity(activity.id)
    assert updated is not None
    assert updated.name == "新活动名称"
    assert updated.status == "completed"

    student_ids = repository.delete_activity(activity.id)
    assert len(student_ids) == 1
    assert repository.get_activity(activity.id) is None


def test_only_required_materials_determine_submission(
    repository: DatabaseRepository, tmp_path: Path
) -> None:
    service = ActivityService(repository)
    activity = service.create(
        "材料清单活动", [("申请表", True), ("补充证明", False)]
    )
    repository.add_students(activity.id, [("张三", None)])
    student = repository.list_students(activity.id)[0]
    requirements = repository.list_activity_requirements(activity.id)
    required = next(item for item in requirements if item.required)
    optional = next(item for item in requirements if not item.required)
    file_service = FileService(repository, tmp_path / "files")
    source = tmp_path / "材料.txt"
    source.write_text("内容", encoding="utf-8")

    assert not student.submitted
    optional_material = file_service.add_files(student.id, [source], optional.id)[0]
    assert optional_material.original_name == "补充证明.txt"
    assert optional_material.stored_path.name == "补充证明.txt"
    assert not repository.list_students(activity.id)[0].submitted
    required_material = file_service.add_files(student.id, [source], required.id)[0]
    assert required_material.original_name == "申请表.txt"
    assert repository.list_students(activity.id)[0].submitted
    export = ExportService(repository).export_activity(activity.id, tmp_path / "output")
    with zipfile.ZipFile(export.exported[0]) as archive:
        assert set(archive.namelist()) == {"申请表.txt", "补充证明.txt"}


def test_duplicate_files_for_requirement_receive_numbered_names(
    repository: DatabaseRepository, tmp_path: Path
) -> None:
    activity = ActivityService(repository).create("材料活动", [("成绩单", True)])
    repository.add_students(activity.id, [("李四", None)])
    student = repository.list_students(activity.id)[0]
    requirement = repository.list_activity_requirements(activity.id)[0]
    source = tmp_path / "任意原名.pdf"
    source.write_bytes(b"pdf")
    service = FileService(repository, tmp_path / "files")

    first = service.add_files(student.id, [source], requirement.id)[0]
    second = service.add_files(student.id, [source], requirement.id)[0]

    assert first.stored_path.name == "成绩单.pdf"
    assert second.stored_path.name == "成绩单_2.pdf"


def test_export_student_statistics_contains_selected_data_only(
    repository: DatabaseRepository, tmp_path: Path
) -> None:
    from openpyxl import load_workbook

    activity = ActivityService(repository).create(
        "统计活动", [("申请表", True), ("补充材料", False)]
    )
    repository.add_students(activity.id, [("张三", "001"), ("李四", "002")])
    students = repository.list_students(activity.id)
    requirements = repository.list_activity_requirements(activity.id)
    required = next(item for item in requirements if item.required)
    source = tmp_path / "原始申请.docx"
    source.write_bytes(b"document")
    target_student = next(item for item in students if item.name == "张三")
    FileService(repository, tmp_path / "files").add_files(
        target_student.id, [source], required.id
    )
    students = repository.list_students(activity.id)

    output = export_student_statistics(
        tmp_path / "统计结果.xlsx",
        students,
        requirements,
        repository,
        set(EXPORT_FIELDS),
    )

    workbook = load_workbook(output, read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    assert rows[0] == (
        "姓名",
        "学号",
        "提交状态",
        "材料进度",
        "文件数量",
        "最后更新时间",
        "申请表",
        "补充材料",
    )
    zhang_san = next(row for row in rows[1:] if row[0] == "张三")
    assert zhang_san[2] == "已提交"
    assert zhang_san[6:] == ("已提交", "未提交")


def test_add_files_updates_status_and_preserves_source(
    repository: DatabaseRepository, tmp_path: Path
) -> None:
    activity = ActivityService(repository).create("材料收集")
    repository.add_students(activity.id, [("张三", "001")])
    student = repository.list_students(activity.id)[0]
    source = tmp_path / "申请书.txt"
    source.write_text("原始内容", encoding="utf-8")
    service = FileService(repository, tmp_path / "files")

    materials = service.add_files(student.id, [source])

    assert source.read_text(encoding="utf-8") == "原始内容"
    assert materials[0].stored_path.read_text(encoding="utf-8") == "原始内容"
    assert repository.list_students(activity.id)[0].submitted

    service.remove_material(materials[0].id)
    assert source.exists()
    assert not materials[0].stored_path.exists()
    assert not repository.list_students(activity.id)[0].submitted


def test_duplicate_source_filenames_do_not_overwrite(
    repository: DatabaseRepository, tmp_path: Path
) -> None:
    activity = ActivityService(repository).create("材料收集")
    repository.add_students(activity.id, [("李四", None)])
    student = repository.list_students(activity.id)[0]
    first_dir = tmp_path / "one"
    second_dir = tmp_path / "two"
    first_dir.mkdir()
    second_dir.mkdir()
    first = first_dir / "材料.txt"
    second = second_dir / "材料.txt"
    first.write_text("第一份", encoding="utf-8")
    second.write_text("第二份", encoding="utf-8")

    materials = FileService(repository, tmp_path / "files").add_files(
        student.id, [first, second]
    )

    assert materials[0].stored_path.name == "材料.txt"
    assert materials[1].stored_path.name == "材料_2.txt"
    assert materials[0].stored_path.read_text(encoding="utf-8") == "第一份"
    assert materials[1].stored_path.read_text(encoding="utf-8") == "第二份"


def test_export_creates_student_zip_with_chinese_filename(
    repository: DatabaseRepository, tmp_path: Path
) -> None:
    activity = ActivityService(repository).create("材料收集")
    repository.add_students(activity.id, [("王五", "0002"), ("赵六", None)])
    student = repository.list_students(activity.id)[0]
    source = tmp_path / "证明材料.txt"
    source.write_text("内容", encoding="utf-8")
    FileService(repository, tmp_path / "files").add_files(student.id, [source])

    result = ExportService(repository).export_activity(activity.id, tmp_path / "output")

    assert len(result.exported) == 1
    assert result.failures == []
    assert result.output_directory == tmp_path / "output" / activity.name
    assert result.exported[0] == result.output_directory / f"{student.name}.zip"
    assert result.exported[0].name == f"{student.name}.zip"
    with zipfile.ZipFile(result.exported[0]) as archive:
        assert archive.namelist() == ["证明材料.txt"]
        assert archive.read("证明材料.txt").decode("utf-8") == "内容"


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("张/三", "张_三"), ("CON", "_CON"), ("...", "未命名")],
)
def test_sanitize_filename(raw: str, expected: str) -> None:
    assert sanitize_filename(raw) == expected
